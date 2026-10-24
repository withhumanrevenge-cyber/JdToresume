"""Part 7: the resume and cover letter as PDF and DOCX, the ATS check on every file, and the download endpoints.

The first half needs no database (real Typst, real python-docx). The API tests use the test database and a fake Groq.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import docx
import pytest
from docx.shared import Mm

from app.features.generation.models import ArtifactRow, GenerationRunRow
from app.features.rendering.ats_roundtrip import check_docx, check_pdf, extract, extract_docx, extractor_agreement
from app.features.rendering.cover_letter import build_cover_letter_document, format_letter_date
from app.features.rendering.docx_renderer import cover_letter_docx, resume_docx
from app.features.rendering.document import RenderOptions, build_resume_document
from app.features.rendering.filenames import content_disposition, download_name
from app.features.rendering.typst_renderer import render_cover_letter_pdf, render_pdf
from app.llm.schemas import CareerProfile, ContactDetails, CoverLetter, JobSpec, TailoredResume
from server.tests.packs import CONTACT, FIXTURES, RESUME, new_application

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _load(model, name):
    return model.model_validate_json((FIXTURES / name).read_text(encoding="utf-8"))


@pytest.fixture
def contact() -> ContactDetails:
    return _load(ContactDetails, "sample_contact.json")


@pytest.fixture
def job() -> JobSpec:
    return _load(JobSpec, "sample_job.json")


@pytest.fixture
def letter() -> CoverLetter:
    return CoverLetter(
        salutation="Dear Hiring Team,",
        paragraphs=["I build payment services in Go; my code is at {{LINK_1}}.", "  ", "{{FULL_NAME}} would love to talk."],
        closing="Best regards,",
        source_fact_ids=[],
        addressed_gap_requirement_ids=[],
    )


# --------------------------------------------------------------------------- cover letter

def test_the_letter_gets_real_details_back_and_a_local_date(letter, contact, job):
    us = build_cover_letter_document(letter, contact, job, region="US", today=date(2026, 10, 21))
    uk = build_cover_letter_document(letter, contact, job, region="UK", today=date(2026, 10, 21))

    assert us["paragraphs"] == ["I build payment services in Go; my code is at github.com/jrivera-example.",
                                "Jordan Rivera would love to talk."]  # placeholders restored, blank paragraph dropped
    assert (us["date"], us["page"]["paper"]) == ("October 21, 2026", "us-letter")
    assert (uk["date"], uk["page"]["paper"]) == ("21 October 2026", "a4")
    assert us["recipient"] == ["Example Pay", "Austin, TX"] and us["signature"] == "Jordan Rivera"
    assert format_letter_date(date(2026, 3, 5), "CA") == "March 5, 2026"


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_the_cover_letter_passes_the_ats_check_in_both_formats(fmt, letter, contact, job):
    document = build_cover_letter_document(letter, contact, job, region="US", today=date(2026, 10, 21))
    if fmt == "pdf":
        check, text = check_pdf(render_cover_letter_pdf(document), contact, [], job), None
    else:
        data = cover_letter_docx(document)
        check, text = check_docx(data, contact, [], job), extract_docx(data).primary
    assert check.parse_ok, check.issues
    assert check.contact_found
    if text:
        assert text.splitlines()[:2] == ["Jordan Rivera", "jordan.rivera@example.com | +1 (555) 010-4477 | Austin, TX | "
                                         "linkedin.com/in/jordan-rivera-example | github.com/jrivera-example"]


# --------------------------------------------------------------------------- resume as DOCX

def test_the_resume_docx_says_what_the_pdf_says(contact, job):
    profile, resume = _load(CareerProfile, "sample_profile.json"), _load(TailoredResume, "sample_tailored_resume.json")
    document = build_resume_document(resume, profile, contact, RenderOptions.for_profile(profile))
    expected = [s["kind"] for s in document["sections"]]

    data = resume_docx(document)
    check = check_docx(data, contact, expected, job)
    pdf_text = extract(render_pdf(document, "classic")).primary

    assert check.parse_ok, check.issues
    assert check.sections_found == expected
    assert extractor_agreement(pdf_text, extract_docx(data).primary) > 0.97  # same words, same order


def test_the_docx_has_a_real_page_size_and_describes_its_owner(contact):
    profile = _load(CareerProfile, "sample_profile.json")
    profile.preferences.region = "UK"
    resume = _load(TailoredResume, "sample_tailored_resume.json")
    document = build_resume_document(resume, profile, contact, RenderOptions.for_profile(profile))

    opened = docx.Document(io.BytesIO(resume_docx(document)))

    section = opened.sections[0]
    assert (round(section.page_width / Mm(1)), round(section.page_height / Mm(1))) == (210, 297)  # A4
    props = opened.core_properties
    assert (props.author, props.title, props.comments) == ("Jordan Rivera", "Jordan Rivera – Resume", "")
    assert [p.text for p in opened.paragraphs if p.style.name == "Heading 1"] == [
        "Summary", "Experience", "Skills", "Projects", "Education", "Certifications"]
    assert not opened.tables and not opened.sections[0].header.paragraphs[0].text  # nothing parsers would skip


def _hostile_docx(*, in_table: bool) -> bytes:
    """Contact details where many parsers don't look: a table, or the page header."""
    doc = docx.Document()
    line = "Jordan Rivera | jordan.rivera@example.com | +1 (555) 010-4477"
    if in_table:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = line
    else:
        doc.sections[0].header.paragraphs[0].text = line
    doc.add_paragraph("Experience", style="Heading 1")
    doc.add_paragraph("Built payment services in Go.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("in_table", [True, False], ids=["table", "page-header"])
def test_the_docx_check_catches_contact_details_parsers_cannot_see(in_table, contact):
    check = check_docx(_hostile_docx(in_table=in_table), contact, ["experience"], None)
    assert not check.parse_ok and not check.contact_found
    if in_table:  # the raw-XML reader sees the table, python-docx's paragraphs don't: the two readers disagree
        assert not check.extractors_agree


# --------------------------------------------------------------------------- file names

@pytest.mark.parametrize(("args", "name", "fallback"), [
    (("Jordan Rivera", "Resume", "Figma", "pdf"), "Jordan-Rivera-Resume-Figma.pdf", "Jordan-Rivera-Resume-Figma.pdf"),
    (("José García", "Cover-Letter", "Société Générale, S.A.", "docx"),
     "José-García-Cover-Letter-Société-Générale-S-A.docx", "Jose-Garcia-Cover-Letter-Societe-Generale-S-A.docx"),
    (("Mary-Jane O'Neil", "Resume", None, "pdf"), "Mary-Jane-ONeil-Resume.pdf", "Mary-Jane-ONeil-Resume.pdf"),
    (("李明", "Resume", "字节跳动", "pdf"), "李明-Resume-字节跳动.pdf", "Resume.pdf"),
    (("Jordan", "Resume", 'x"; filename=evil.exe', "pdf"), "Jordan-Resume-x-filename-evil-exe.pdf",
     "Jordan-Resume-x-filename-evil-exe.pdf"),
], ids=["plain", "accents", "apostrophe-no-company", "non-latin", "header-injection"])
def test_download_names(args, name, fallback):
    assert download_name(*args) == name
    header = content_disposition(name)
    assert header.startswith(f'attachment; filename="{fallback}"; filename*=UTF-8\'\'')
    assert header.count('"') == 2  # nothing from the input can close the quoted name early


def test_a_long_company_name_is_shortened_but_the_document_label_survives():
    assert download_name("Jordan Rivera", "Resume", "A" * 90, "pdf") == f"Jordan-Rivera-Resume-{'A' * 40}.pdf"


# --------------------------------------------------------------------------- the API

async def _generated(client, headers, *, with_contact: bool = True) -> str:
    application_id = await new_application(client, headers)
    if with_contact:
        assert (await client.put("/v1/profile/contact", headers=headers, json=CONTACT)).status_code == 200
    assert (await client.post(f"/v1/applications/{application_id}/generate", headers=headers)).status_code == 200
    return application_id


@pytest.mark.anyio
@pytest.mark.parametrize(("file_name", "media_type", "magic", "expected_name"), [
    ("resume.pdf", "application/pdf", b"%PDF", "Jordan-Rivera-Resume-Example-Pay.pdf"),
    ("resume.docx", DOCX_TYPE, b"PK", "Jordan-Rivera-Resume-Example-Pay.docx"),
    ("cover-letter.pdf", "application/pdf", b"%PDF", "Jordan-Rivera-Cover-Letter-Example-Pay.pdf"),
    ("cover-letter.docx", DOCX_TYPE, b"PK", "Jordan-Rivera-Cover-Letter-Example-Pay.docx"),
])
async def test_every_file_downloads_checked_and_properly_named(client, sign_in, ai, file_name, media_type, magic,
                                                               expected_name):
    headers = sign_in("jordan@example.com")
    application_id = await _generated(client, headers)

    response = await client.get(f"/v1/applications/{application_id}/files/{file_name}", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == media_type and response.content.startswith(magic)
    assert f'filename="{expected_name}"' in response.headers["content-disposition"]
    assert response.headers["x-ats-check"] == "passed"
    assert response.headers["cache-control"] == "private, no-store"


@pytest.mark.anyio
async def test_what_the_ats_sees(client, sign_in, ai):
    headers = sign_in("jordan@example.com")
    application_id = await _generated(client, headers)

    pdf = (await client.get(f"/v1/applications/{application_id}/ats-check", headers=headers)).json()
    word = (await client.get(f"/v1/applications/{application_id}/ats-check?document=cover_letter&format=docx",
                             headers=headers)).json()

    assert pdf["check"]["parse_ok"] and pdf["pages"] == 1 and pdf["file_name"] == "Jordan-Rivera-Resume-Example-Pay.pdf"
    assert pdf["text"].startswith("Jordan Rivera\n") and "\nExperience\n" in pdf["text"]
    assert pdf["check"]["keyword_coverage"] is not None  # measured against the job, in the extracted text
    assert (word["document"], word["format"], word["pages"]) == ("cover_letter", "docx", None)
    assert word["check"]["parse_ok"] and "Dear Hiring Team," in word["text"]


@pytest.mark.anyio
async def test_downloads_need_a_generated_pack_and_contact_details(client, sign_in, ai):
    headers = sign_in("jordan@example.com")
    fresh = await new_application(client, headers)
    not_yet = await client.get(f"/v1/applications/{fresh}/files/resume.pdf", headers=headers)
    assert (not_yet.status_code, not_yet.json()["code"]) == (404, "not_generated")

    no_contact = await _generated(client, sign_in("priya@example.com"), with_contact=False)
    response = await client.get(f"/v1/applications/{no_contact}/files/resume.pdf", headers=sign_in("priya@example.com"))
    assert (response.status_code, response.json()["code"]) == (422, "contact_required")


@pytest.mark.anyio
async def test_nobody_can_download_someone_elses_documents(client, sign_in, ai):
    application_id = await _generated(client, sign_in("jordan@example.com"))
    stranger = sign_in("priya@example.com")
    for path in ("files/resume.pdf", "ats-check"):
        response = await client.get(f"/v1/applications/{application_id}/{path}", headers=stranger)
        assert (response.status_code, response.json()["code"]) == (404, "not_found")


@pytest.mark.anyio
async def test_unknown_file_names_and_templates_are_rejected(client, sign_in, ai):
    headers = sign_in("jordan@example.com")
    application_id = await _generated(client, headers)
    for path in ("files/resume.exe", "files/resume.pdf?template=fancy", "ats-check?format=txt"):
        assert (await client.get(f"/v1/applications/{application_id}/{path}", headers=headers)).status_code == 422


@pytest.mark.anyio
async def test_the_resume_is_rendered_from_the_profile_version_it_was_written_from(client, sign_in, ai):
    headers = sign_in("jordan@example.com")
    application_id = await _generated(client, headers)
    profile = (await client.get("/v1/profile", headers=headers)).json()
    profile["profile"]["experience"][0]["title"] = "Staff Engineer"  # edited after the pack was written
    saved = await client.put("/v1/profile", headers=headers,
                             json={"base_version": profile["version"], "profile": profile["profile"]})
    assert saved.status_code == 200

    text = (await client.get(f"/v1/applications/{application_id}/ats-check", headers=headers)).json()["text"]

    assert "Software Engineer II, Northwind Payments" in text and "Staff Engineer" not in text


@pytest.mark.anyio
async def test_a_failing_check_is_reported_and_a_later_partial_run_hides_nothing(client, sign_in, ai, db):
    headers = sign_in("jordan@example.com")
    application_id = await _generated(client, headers)
    # A newer run that saved only a resume, with the oldest job listed first (an ATS reads that as out of order).
    oldest_first = {**RESUME, "experience": list(reversed(RESUME["experience"]))}
    async with db.sessions() as session:
        run = GenerationRunRow(application_id=uuid.UUID(application_id), profile_version=1, mode="deep", status="failed")
        session.add(run)
        await session.flush()
        session.add(ArtifactRow(run_id=run.id, kind="resume", content={"resume": oldest_first}))
        await session.commit()

    resume = await client.get(f"/v1/applications/{application_id}/files/resume.pdf", headers=headers)
    view = (await client.get(f"/v1/applications/{application_id}/ats-check", headers=headers)).json()
    letter = await client.get(f"/v1/applications/{application_id}/files/cover-letter.pdf", headers=headers)

    assert resume.status_code == 200 and resume.headers["x-ats-check"] == "failed"  # delivered, not hidden
    assert any("reverse-chronological" in issue for issue in view["check"]["issues"])
    assert letter.status_code == 200 and letter.headers["x-ats-check"] == "passed"  # from the earlier run
